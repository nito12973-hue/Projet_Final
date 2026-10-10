"""
Gestion des paiements : liste, export CSV, règlement manuel et bordereaux de télétransmission B2B.
"""

import csv
from decimal import Decimal

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db.models import Q, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..forms import PaiementReglementForm
from ..models import Delivrance, JournalActivite, Paiement, Prestataire, User
from .utils import _cellule_csv, _paginer, _trier, admin_required, journaliser


def _filtrer_paiements(request):
    """Filtres partages entre la liste et l'export CSV des paiements."""
    paiements = Paiement.objects.select_related(
        "consultation", "consultation__patient", "consultation__service"
    )

    statut = request.GET.get("statut", "")
    if statut:
        paiements = paiements.filter(statut=statut)

    recherche = request.GET.get("q", "").strip()
    if recherche:
        paiements = paiements.filter(
            Q(consultation__patient__nom__icontains=recherche)
            | Q(consultation__patient__prenom__icontains=recherche)
        )

    paiements = _trier(
        request, paiements,
        ["consultation__patient__nom", "montant_total", "montant_part_assurance", "montant_part_patient", "statut"],
        "-consultation__date_consultation",
    )

    return paiements, {"statut": statut, "recherche": recherche}


@admin_required
def liste_paiements(request):
    paiements, filtres = _filtrer_paiements(request)

    totaux = Paiement.objects.aggregate(
        total_regle=Sum("montant_part_patient", filter=Q(statut=Paiement.Statut.REGLE)),
        total_non_regle=Sum("montant_part_patient", filter=Q(statut=Paiement.Statut.NON_REGLE)),
    )

    contexte = {
        "paiements": _paginer(request, paiements),
        "statut_choisi": filtres["statut"],
        "statuts": Paiement.Statut.choices,
        "recherche": filtres["recherche"],
        "total_regle": totaux["total_regle"] or 0,
        "total_non_regle": totaux["total_non_regle"] or 0,
    }
    return render(request, "liste_paiements.html", contexte)


@admin_required
def exporter_paiements_csv(request):
    paiements, _ = _filtrer_paiements(request)

    reponse = HttpResponse(content_type="text/csv")
    reponse["Content-Disposition"] = 'attachment; filename="paiements_santesn.csv"'
    reponse.write("\ufeff")  # BOM : Excel (FR) detecte l'UTF-8 sans le confondre avec l'encodage local.
    ecrivain = csv.writer(reponse, delimiter=";")
    ecrivain.writerow([
        "Reference", "Patient", "Date de consultation", "Montant total",
        "Part assurance", "Part patient", "Statut", "Mode de reglement",
        "Date de reglement",
    ])
    for paiement in paiements:
        ecrivain.writerow([_cellule_csv(v) for v in [
            paiement.pk,
            str(paiement.consultation.patient),
            paiement.consultation.date_consultation.strftime("%d/%m/%Y %H:%M"),
            paiement.montant_total,
            paiement.montant_part_assurance,
            paiement.montant_part_patient,
            paiement.get_statut_display(),
            paiement.get_mode_reglement_display() if paiement.mode_reglement else "",
            paiement.date_reglement.strftime("%d/%m/%Y %H:%M") if paiement.date_reglement else "",
        ]])
    return reponse


@admin_required
def marquer_paiement_regle(request, pk):
    paiement = get_object_or_404(Paiement, pk=pk)
    if paiement.statut == Paiement.Statut.REGLE:
        messages.info(request, "Ce paiement est déjà réglé.")
        return redirect("liste_paiements")

    if request.method == "POST":
        form = PaiementReglementForm(request.POST, instance=paiement)
        if form.is_valid():
            paiement = form.save(commit=False)
            paiement.statut = Paiement.Statut.REGLE
            paiement.date_reglement = timezone.now()
            # Qui a constate l'encaissement : jamais saisi, toujours deduit
            # du compte connecte. En especes, c'est la seule trace de la
            # personne qui a recu l'argent.
            paiement.enregistre_par = request.user
            paiement.save()
            journaliser(
                request, JournalActivite.Action.REGLEMENT,
                f"Paiement #{paiement.pk} · {paiement.consultation.patient}",
                f"{paiement.montant_part_patient} F CFA · {paiement.get_mode_reglement_display()}",
            )
            messages.success(request, f"Le règlement pour {paiement.consultation.patient} a été enregistré avec succès.")
            return redirect("liste_paiements")
    else:
        form = PaiementReglementForm(instance=paiement)

    return render(request, "marquer_paiement_regle.html", {
        "paiement": paiement,
        "form": form,
    })


@login_required
def recu_paiement(request, pk):
    """Reçu / Quittance officielle de paiement au format A5/Ticket."""
    paiement = get_object_or_404(
        Paiement.objects.select_related(
            "consultation__patient",
            "consultation__medecin",
            "consultation__service",
            "consultation__prise_en_charge",
            "consultation__patient__plan_couverture",
            "consultation__patient__assure_principal__plan_couverture",
            "enregistre_par",
        ),
        pk=pk,
    )

    est_admin = request.user.role == User.Role.ADMIN
    est_medecin = (
        request.user.role == User.Role.MEDECIN
        and hasattr(request.user, "medecin")
        and paiement.consultation.medecin == request.user.medecin
    )
    est_assure = False
    if request.user.role == User.Role.ASSURE and hasattr(request.user, "patient"):
        patient_assure = request.user.patient
        membres = [patient_assure.pk] + list(patient_assure.ayants_droit.values_list("pk", flat=True))
        if paiement.consultation.patient_id in membres:
            est_assure = True

    if not (est_admin or est_medecin or est_assure):
        raise PermissionDenied("Vous n'avez pas accès à ce reçu de paiement.")

    date_ref = paiement.date_reglement or paiement.consultation.date_consultation
    reference_recu = f"REC-{date_ref.strftime('%Y%m%d')}-{paiement.pk:05d}"

    return render(request, "recu_paiement.html", {
        "paiement": paiement,
        "reference_recu": reference_recu,
    })


# ---------------------------------------------------------------------------
# Module B2B : Télétransmission & Bordereaux Mensuels de Tiers Payant IPM
# ---------------------------------------------------------------------------

class NumberedCanvasLandscape:
    """Générateur de Canvas ReportLab Paysage avec numérotation 'Page X sur Y' et en-tête institutionnel."""

    def __init__(self, *args, **kwargs):
        from reportlab.pdfgen import canvas
        self._canvas_class = canvas.Canvas
        self._underlying = canvas.Canvas(*args, **kwargs)
        self._saved_page_states = []

    def __getattr__(self, name):
        return getattr(self._underlying, name)

    def showPage(self):
        self._saved_page_states.append(dict(self._underlying.__dict__))
        self._underlying._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self._underlying.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            self._underlying.showPage()
        self._underlying.save()

    def draw_page_decorations(self, page_count):
        self._underlying.saveState()

        # En-tête institutionnel
        self._underlying.setFont("Helvetica-Bold", 8.5)
        self._underlying.setFillColor(colors.HexColor("#0e7c86"))
        self._underlying.drawString(40, 568, "SANTÉSN — SYSTÈME NATIONAL DE TIERS-PAYANT")
        self._underlying.setFont("Helvetica", 7.5)
        self._underlying.setFillColor(colors.HexColor("#5f6f7d"))
        self._underlying.drawString(290, 568, "Bordereau Récapitulatif de Télétransmission & Facturation Mensuelle IPM")

        self._underlying.setStrokeColor(colors.HexColor("#cbd5e1"))
        self._underlying.setLineWidth(0.5)
        self._underlying.line(40, 560, 802, 560)

        # Pied de page institutionnel
        self._underlying.line(40, 36, 802, 36)
        self._underlying.setFont("Helvetica", 7.5)
        self._underlying.setFillColor(colors.HexColor("#5f6f7d"))
        self._underlying.drawString(40, 24, "Document officiel de télétransmission certifié conforme aux conventions IPM en vigueur.")
        texte_page = f"Page {self._underlying._pageNumber} sur {page_count}"
        self._underlying.drawRightString(802, 24, texte_page)

        self._underlying.restoreState()


LISTE_MOIS = [
    (1, "Janvier"), (2, "Février"), (3, "Mars"), (4, "Avril"),
    (5, "Mai"), (6, "Juin"), (7, "Juillet"), (8, "Août"),
    (9, "Septembre"), (10, "Octobre"), (11, "Novembre"), (12, "Décembre"),
]


def _verifier_acces_teletransmission(request):
    """Vérifie que l'utilisateur a accès au bordereau (Admin, Médecin, Pharmacien)."""
    if not request.user.is_authenticated:
        raise PermissionDenied("Connexion requise.")
    if request.user.role not in (User.Role.ADMIN, User.Role.MEDECIN, User.Role.PHARMACIEN):
        raise PermissionDenied("Accès réservé aux gestionnaires IPM et professionnels de santé.")

    prestataire_force = None
    if request.user.role == User.Role.PHARMACIEN and hasattr(request.user, "pharmacien"):
        prestataire_force = request.user.pharmacien.prestataire
    elif request.user.role == User.Role.MEDECIN and hasattr(request.user, "medecin"):
        prestataire_force = request.user.medecin.prestataire

    est_admin = request.user.role == User.Role.ADMIN
    return est_admin, prestataire_force


def _recuperer_actes_teletransmission(annee=None, mois=None, prestataire_id=None, type_acte="tous", recherche=""):
    """
    Consolide l'ensemble des actes de Tiers Payant (Consultations médicales + Délivrances pharmacie)
    éligibles au remboursement par l'IPM (part assurance > 0).
    """
    maintenant = timezone.now()
    if annee is None:
        annee = maintenant.year

    actes = []

    # 1. Consultations Médicales couvertes par l'assurance
    if type_acte in ("tous", "consultation"):
        paiements_qs = Paiement.objects.select_related(
            "consultation__patient",
            "consultation__medecin__prestataire",
            "consultation__service",
            "consultation__prise_en_charge",
        ).filter(montant_part_assurance__gt=Decimal("0.00"))

        if annee:
            paiements_qs = paiements_qs.filter(consultation__date_consultation__year=annee)
        if mois:
            paiements_qs = paiements_qs.filter(consultation__date_consultation__month=mois)
        if prestataire_id:
            paiements_qs = paiements_qs.filter(consultation__medecin__prestataire_id=prestataire_id)
        if recherche:
            paiements_qs = paiements_qs.filter(
                Q(consultation__patient__nom__icontains=recherche)
                | Q(consultation__patient__prenom__icontains=recherche)
                | Q(consultation__patient__numero_carte__icontains=recherche)
            )

        for p in paiements_qs:
            medecin = p.consultation.medecin
            prestataire = medecin.prestataire if medecin else None
            patient = p.consultation.patient
            actes.append({
                "id": f"CS-{p.pk}",
                "type": "Consultation",
                "type_code": "consultation",
                "reference": f"CS-{p.consultation.pk:05d}",
                "date": p.consultation.date_consultation,
                "patient_nom": f"{patient.prenom} {patient.nom}",
                "patient_carte": patient.numero_carte,
                "praticien": str(medecin) if medecin else "Médecin",
                "prestataire": str(prestataire) if prestataire else "Cabinet Médical",
                "prestataire_id": prestataire.pk if prestataire else None,
                "description": p.consultation.service.nom if p.consultation.service else "Consultation médicale",
                "montant_total": p.montant_total or Decimal("0.00"),
                "taux": p.taux_applique or Decimal("0.00"),
                "part_assurance": p.montant_part_assurance or Decimal("0.00"),
                "part_patient": p.montant_part_patient or Decimal("0.00"),
            })

    # 2. Délivrances Pharmaceutiques
    if type_acte in ("tous", "pharmacie"):
        delivrances_qs = Delivrance.objects.select_related(
            "ordonnance__consultation__patient",
            "ordonnance__consultation__medecin",
            "pharmacien__prestataire",
        ).filter(montant_part_assurance__gt=Decimal("0.00"))

        if annee:
            delivrances_qs = delivrances_qs.filter(date_delivrance__year=annee)
        if mois:
            delivrances_qs = delivrances_qs.filter(date_delivrance__month=mois)
        if prestataire_id:
            delivrances_qs = delivrances_qs.filter(pharmacien__prestataire_id=prestataire_id)
        if recherche:
            delivrances_qs = delivrances_qs.filter(
                Q(ordonnance__consultation__patient__nom__icontains=recherche)
                | Q(ordonnance__consultation__patient__prenom__icontains=recherche)
                | Q(ordonnance__consultation__patient__numero_carte__icontains=recherche)
                | Q(ordonnance__code_qr__icontains=recherche)
            )

        for d in delivrances_qs:
            patient = d.ordonnance.consultation.patient
            pharmacien = d.pharmacien
            prestataire = pharmacien.prestataire if pharmacien else None
            actes.append({
                "id": f"DEL-{d.pk}",
                "type": "Pharmacie",
                "type_code": "pharmacie",
                "reference": f"DEL-{d.ordonnance.code_qr}",
                "date": d.date_delivrance,
                "patient_nom": f"{patient.prenom} {patient.nom}",
                "patient_carte": patient.numero_carte,
                "praticien": str(pharmacien) if pharmacien else "Pharmacien",
                "prestataire": str(prestataire) if prestataire else "Pharmacie Partenaire",
                "prestataire_id": prestataire.pk if prestataire else None,
                "description": f"Délivrance Rx #{d.ordonnance.code_qr}",
                "montant_total": d.montant_total or Decimal("0.00"),
                "taux": d.taux_couverture or Decimal("0.00"),
                "part_assurance": d.montant_part_assurance or Decimal("0.00"),
                "part_patient": d.montant_part_patient or Decimal("0.00"),
            })

    # Tri par date décroissante
    actes.sort(key=lambda a: a["date"], reverse=True)

    # Calcul des totaux consolidés
    total_brut = sum((a["montant_total"] for a in actes), Decimal("0.00"))
    total_part_assurance = sum((a["part_assurance"] for a in actes), Decimal("0.00"))
    total_part_patient = sum((a["part_patient"] for a in actes), Decimal("0.00"))

    totaux = {
        "nb_actes": len(actes),
        "total_brut": total_brut,
        "total_part_assurance": total_part_assurance,
        "total_part_patient": total_part_patient,
    }
    return actes, totaux


@login_required
def bordereau_teletransmission(request):
    """Tableau de bord de télétransmission et consultation du bordereau mensuel."""
    est_admin, prestataire_force = _verifier_acces_teletransmission(request)

    maintenant = timezone.now()
    try:
        annee = int(request.GET.get("annee", maintenant.year))
    except (TypeError, ValueError):
        annee = maintenant.year

    mois_raw = request.GET.get("mois", "")
    try:
        mois = int(mois_raw) if mois_raw else maintenant.month
    except (TypeError, ValueError):
        mois = maintenant.month

    if mois == 0:
        mois = None

    type_acte = request.GET.get("type_acte", "tous")
    if type_acte not in ("tous", "consultation", "pharmacie"):
        type_acte = "tous"

    recherche = request.GET.get("q", "").strip()

    # Si praticien / pharmacien restreint, il ne peut voir que sa structure
    if prestataire_force:
        prestataire_id = prestataire_force.pk
    else:
        prestataire_id_raw = request.GET.get("prestataire", "")
        try:
            prestataire_id = int(prestataire_id_raw) if prestataire_id_raw else None
        except (TypeError, ValueError):
            prestataire_id = None

    actes, totaux = _recuperer_actes_teletransmission(
        annee=annee,
        mois=mois,
        prestataire_id=prestataire_id,
        type_acte=type_acte,
        recherche=recherche,
    )

    liste_annees = list(range(maintenant.year - 2, maintenant.year + 2))
    prestataires = Prestataire.objects.filter(partenaire=True).order_by("nom") if est_admin else []

    # Nom du prestataire sélectionné
    nom_prestataire = "Tous les prestataires conventionnés"
    if prestataire_force:
        nom_prestataire = str(prestataire_force)
    elif prestataire_id:
        p_obj = Prestataire.objects.filter(pk=prestataire_id).first()
        if p_obj:
            nom_prestataire = str(p_obj)

    contexte = {
        "actes": _paginer(request, actes),
        "totaux": totaux,
        "liste_mois": LISTE_MOIS,
        "liste_annees": liste_annees,
        "prestataires": prestataires,
        "annee_choisie": annee,
        "mois_choisi": mois or 0,
        "type_choisi": type_acte,
        "prestataire_choisi": prestataire_id or "",
        "nom_prestataire": nom_prestataire,
        "prestataire_force": prestataire_force,
        "est_admin": est_admin,
        "recherche": recherche,
    }
    return render(request, "bordereau_teletransmission.html", contexte)


@login_required
def exporter_bordereau_teletransmission_pdf(request):
    """Génère le bordereau officiel de télétransmission au format PDF paysage certifié."""
    est_admin, prestataire_force = _verifier_acces_teletransmission(request)

    maintenant = timezone.now()
    try:
        annee = int(request.GET.get("annee", maintenant.year))
    except (TypeError, ValueError):
        annee = maintenant.year

    mois_raw = request.GET.get("mois", "")
    try:
        mois = int(mois_raw) if mois_raw else maintenant.month
    except (TypeError, ValueError):
        mois = maintenant.month

    if mois == 0:
        mois = None

    type_acte = request.GET.get("type_acte", "tous")
    recherche = request.GET.get("q", "").strip()

    if prestataire_force:
        prestataire_id = prestataire_force.pk
    else:
        prestataire_id_raw = request.GET.get("prestataire", "")
        try:
            prestataire_id = int(prestataire_id_raw) if prestataire_id_raw else None
        except (TypeError, ValueError):
            prestataire_id = None

    actes, totaux = _recuperer_actes_teletransmission(
        annee=annee,
        mois=mois,
        prestataire_id=prestataire_id,
        type_acte=type_acte,
        recherche=recherche,
    )

    nom_prestataire = "Tous les prestataires conventionnés"
    if prestataire_force:
        nom_prestataire = str(prestataire_force)
    elif prestataire_id:
        p_obj = Prestataire.objects.filter(pk=prestataire_id).first()
        if p_obj:
            nom_prestataire = str(p_obj)

    libelle_mois = "Année complète"
    if mois:
        dict_mois = dict(LISTE_MOIS)
        libelle_mois = f"{dict_mois.get(mois, '')} {annee}"
    else:
        libelle_mois = f"Exercice {annee}"

    reponse = HttpResponse(content_type="application/pdf")
    nom_fichier = f"bordereau_teletransmission_{annee}_{mois or 'annuel'}.pdf"
    reponse["Content-Disposition"] = f'attachment; filename="{nom_fichier}"'

    document = SimpleDocTemplate(
        reponse,
        pagesize=landscape(A4),
        leftMargin=40,
        rightMargin=40,
        topMargin=45,
        bottomMargin=45,
    )

    styles = getSampleStyleSheet()

    style_titre = ParagraphStyle(
        "TitreBordereau",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#0b2027"),
        alignment=0,
    )
    style_sous_titre = ParagraphStyle(
        "SousTitreBordereau",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#0e7c86"),
    )
    style_meta = ParagraphStyle(
        "MetaBordereau",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#5f6f7d"),
    )
    style_cell = ParagraphStyle(
        "CellNormal",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#1e293b"),
    )
    style_cell_bold = ParagraphStyle(
        "CellBold",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#0f172a"),
    )
    style_cell_center = ParagraphStyle(
        "CellCenter",
        parent=style_cell,
        alignment=1,
    )
    style_cell_right = ParagraphStyle(
        "CellRight",
        parent=style_cell,
        alignment=2,
    )
    style_cell_right_bold = ParagraphStyle(
        "CellRightBold",
        parent=style_cell_bold,
        alignment=2,
    )

    elements = [
        Paragraph("BORDEREAU RÉCAPITULATIF DE TÉLÉTRANSMISSION TIERS PAYANT", style_titre),
        Paragraph(f"Établissement / Prestataire : <b>{nom_prestataire}</b> · Période : <b>{libelle_mois}</b>", style_sous_titre),
        Paragraph(
            f"Généré le {maintenant.strftime('%d/%m/%Y à %H:%M')} par {request.user.get_full_name() or request.user.email} · "
            f"Actes transmis : <b>{totaux['nb_actes']}</b> · "
            f"Total Tiers-Payant Réclamé : <b>{totaux['total_part_assurance']:,.0f} FCFA</b>",
            style_meta,
        ),
        Spacer(1, 10),
    ]

    # Tableau Récapitulatif Financier (Synthèse du bordereau)
    table_kpi_data = [
        [
            Paragraph(f"<font size=6.5 color='#5f6f7d'>NOMBRE D'ACTES</font><br/><font size=11 color='#095059'><b>{totaux['nb_actes']}</b></font>", styles["Normal"]),
            Paragraph(f"<font size=6.5 color='#5f6f7d'>MONTANT TOTAL BRUT</font><br/><font size=11 color='#095059'><b>{totaux['total_brut']:,.0f} FCFA</b></font>", styles["Normal"]),
            Paragraph(f"<font size=6.5 color='#5f6f7d'>PART PATIENTS (TICKET MODÉRATEUR)</font><br/><font size=11 color='#095059'><b>{totaux['total_part_patient']:,.0f} FCFA</b></font>", styles["Normal"]),
            Paragraph(f"<font size=6.5 color='#5f6f7d'><b>PART IPM RÉCLAMÉE (NET À REMBOURSER)</b></font><br/><font size=11 color='#0e7c86'><b>{totaux['total_part_assurance']:,.0f} FCFA</b></font>", styles["Normal"]),
        ]
    ]
    table_kpi = Table(table_kpi_data, colWidths=[130, 200, 210, 221])
    table_kpi.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ('BOX', (0, 0), (-1, -1), 0.75, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
        ('RIGHTPADDING', (0, 0), (-1, -1), 10),
    ]))
    elements.append(table_kpi)
    elements.append(Spacer(1, 12))

    # Tableau Détaillé des Actes
    entetes = [
        Paragraph("<b>Réf Acte</b>", style_cell_bold),
        Paragraph("<b>Date</b>", style_cell_bold),
        Paragraph("<b>Type</b>", style_cell_bold),
        Paragraph("<b>Patient & Carte IPM</b>", style_cell_bold),
        Paragraph("<b>Praticien / Structure</b>", style_cell_bold),
        Paragraph("<b>Prestation</b>", style_cell_bold),
        Paragraph("<b>Total Brut</b>", style_cell_right_bold),
        Paragraph("<b>Taux</b>", style_cell_center),
        Paragraph("<b>Part IPM</b>", style_cell_right_bold),
        Paragraph("<b>Ticket Mod.</b>", style_cell_right_bold),
    ]

    col_widths = [65, 55, 60, 130, 125, 120, 60, 35, 55, 55]
    table_rows = [entetes]

    for a in actes:
        table_rows.append([
            Paragraph(f"<font color='#0e7c86'><b>{a['reference']}</b></font>", style_cell),
            Paragraph(a["date"].strftime("%d/%m/%Y"), style_cell),
            Paragraph(a["type"], style_cell),
            Paragraph(f"<b>{a['patient_nom']}</b><br/><font size=6.5 color='#5f6f7d'>{a['patient_carte']}</font>", style_cell),
            Paragraph(f"{a['praticien']}<br/><font size=6.5 color='#5f6f7d'>{a['prestataire']}</font>", style_cell),
            Paragraph(a["description"], style_cell),
            Paragraph(f"{a['montant_total']:,.0f}", style_cell_right),
            Paragraph(f"{a['taux']:.0f}%", style_cell_center),
            Paragraph(f"<b>{a['part_assurance']:,.0f}</b>", style_cell_right_bold),
            Paragraph(f"{a['part_patient']:,.0f}", style_cell_right),
        ])

    # Ligne de total final
    table_rows.append([
        Paragraph("<b>TOTAL</b>", style_cell_bold),
        Paragraph("", style_cell),
        Paragraph("", style_cell),
        Paragraph(f"<b>{totaux['nb_actes']} actes</b>", style_cell_bold),
        Paragraph("", style_cell),
        Paragraph("", style_cell),
        Paragraph(f"<b>{totaux['total_brut']:,.0f}</b>", style_cell_right_bold),
        Paragraph("", style_cell),
        Paragraph(f"<b>{totaux['total_part_assurance']:,.0f}</b>", style_cell_right_bold),
        Paragraph(f"<b>{totaux['total_part_patient']:,.0f}</b>", style_cell_right_bold),
    ])

    table_actes = Table(table_rows, colWidths=col_widths, repeatRows=1)
    table_actes.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0e7c86")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#f1f5f9")),
    ]))

    # Réadapter le style des entêtes blanches
    for col_idx in range(len(entetes)):
        entete_p = entetes[col_idx]
        entete_p.style.textColor = colors.white

    elements.append(table_actes)
    elements.append(Spacer(1, 15))

    # Bloc de signatures et attestations légales
    signatures_data = [
        [
            Paragraph(
                "<b>POUR LE PRESTATAIRE CONVENTIONNÉ</b><br/><br/>"
                "Je certifie l'exactitude des soins et fournitures dispensés aux assurés ci-dessus désignés.<br/><br/>"
                "Date : ........................................<br/>"
                "Cachet & Signature du Directeur / Pharmacien :",
                styles["Normal"],
            ),
            Paragraph(
                "<b>POUR L'IPM / L'ORGANISME ASSUREUR</b><br/><br/>"
                "Contrôle médical et liquidation validés. Bon à payer pour le montant certifié ci-contre.<br/><br/>"
                "Date : ........................................<br/>"
                "Visa du Contrôleur & Signature de l'Ordonnateur :",
                styles["Normal"],
            ),
        ]
    ]
    table_sig = Table(signatures_data, colWidths=[375, 386])
    table_sig.setStyle(TableStyle([
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#94a3b8")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 16),
        ('LEFTPADDING', (0, 0), (-1, -1), 12),
        ('RIGHTPADDING', (0, 0), (-1, -1), 12),
    ]))
    elements.append(KeepTogether(table_sig))

    document.build(elements, canvasmaker=NumberedCanvasLandscape)
    return reponse


@login_required
def exporter_bordereau_teletransmission_excel(request):
    """Exporte le bordereau de télétransmission sous format Excel professionnel (.xlsx)."""
    est_admin, prestataire_force = _verifier_acces_teletransmission(request)

    maintenant = timezone.now()
    try:
        annee = int(request.GET.get("annee", maintenant.year))
    except (TypeError, ValueError):
        annee = maintenant.year

    mois_raw = request.GET.get("mois", "")
    try:
        mois = int(mois_raw) if mois_raw else maintenant.month
    except (TypeError, ValueError):
        mois = maintenant.month

    if mois == 0:
        mois = None

    type_acte = request.GET.get("type_acte", "tous")
    recherche = request.GET.get("q", "").strip()

    if prestataire_force:
        prestataire_id = prestataire_force.pk
    else:
        prestataire_id_raw = request.GET.get("prestataire", "")
        try:
            prestataire_id = int(prestataire_id_raw) if prestataire_id_raw else None
        except (TypeError, ValueError):
            prestataire_id = None

    actes, totaux = _recuperer_actes_teletransmission(
        annee=annee,
        mois=mois,
        prestataire_id=prestataire_id,
        type_acte=type_acte,
        recherche=recherche,
    )

    nom_prestataire = "Tous les prestataires conventionnés"
    if prestataire_force:
        nom_prestataire = str(prestataire_force)
    elif prestataire_id:
        p_obj = Prestataire.objects.filter(pk=prestataire_id).first()
        if p_obj:
            nom_prestataire = str(p_obj)

    libelle_mois = "Année complète"
    if mois:
        dict_mois = dict(LISTE_MOIS)
        libelle_mois = f"{dict_mois.get(mois, '')} {annee}"
    else:
        libelle_mois = f"Exercice {annee}"

    classeur = openpyxl.Workbook()
    ws = classeur.active
    ws.title = "Bordereau Tiers-Payant"
    ws.views.sheetView[0].showGridLines = True

    fill_entete = PatternFill(start_color="0E7C86", end_color="0E7C86", fill_type="solid")
    fill_total = PatternFill(start_color="E2F1F3", end_color="E2F1F3", fill_type="solid")
    font_entete = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_titre = Font(name="Calibri", size=13, bold=True, color="0B2027")
    font_sous_titre = Font(name="Calibri", size=10, italic=True, color="5F6F7D")
    font_bold = Font(name="Calibri", size=10, bold=True, color="0B2027")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    border_fine = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )

    # Titre et métadonnées
    ws["A1"] = "SantéSN — Bordereau Récapitulatif de Télétransmission Tiers-Payant"
    ws["A1"].font = font_titre
    ws["A2"] = f"Établissement : {nom_prestataire} · Période : {libelle_mois}"
    ws["A2"].font = font_sous_titre
    ws["A3"] = f"Généré le {maintenant.strftime('%d/%m/%Y à %H:%M')} par {request.user.get_full_name() or request.user.email} · Actes : {totaux['nb_actes']}"
    ws["A3"].font = font_sous_titre

    # En-têtes de colonnes
    headers = [
        "N° Référence", "Date de l'acte", "Type d'acte", "Bénéficiaire (Patient)",
        "N° Carte Assuré", "Praticien / Intervenant", "Établissement / Structure",
        "Prestation / Médicaments", "Montant Brut (FCFA)", "Taux IPM (%)",
        "Part IPM Réclamée (FCFA)", "Ticket Modérateur Patient (FCFA)",
    ]

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=header)
        cell.fill = fill_entete
        cell.font = font_entete
        cell.alignment = align_center
        cell.border = border_fine

    # Données
    current_row = 6
    for a in actes:
        ws.cell(row=current_row, column=1, value=a["reference"]).alignment = align_center
        ws.cell(row=current_row, column=2, value=a["date"].strftime("%d/%m/%Y %H:%M")).alignment = align_center
        ws.cell(row=current_row, column=3, value=a["type"]).alignment = align_center
        ws.cell(row=current_row, column=4, value=a["patient_nom"])
        ws.cell(row=current_row, column=5, value=a["patient_carte"]).alignment = align_center
        ws.cell(row=current_row, column=6, value=a["praticien"])
        ws.cell(row=current_row, column=7, value=a["prestataire"])
        ws.cell(row=current_row, column=8, value=a["description"])

        c_brut = ws.cell(row=current_row, column=9, value=float(a["montant_total"]))
        c_brut.number_format = '#,##0 "FCFA"'
        c_brut.alignment = align_right

        c_taux = ws.cell(row=current_row, column=10, value=float(a["taux"]))
        c_taux.number_format = '0"%"'
        c_taux.alignment = align_center

        c_ass = ws.cell(row=current_row, column=11, value=float(a["part_assurance"]))
        c_ass.number_format = '#,##0 "FCFA"'
        c_ass.alignment = align_right
        c_ass.font = font_bold

        c_pat = ws.cell(row=current_row, column=12, value=float(a["part_patient"]))
        c_pat.number_format = '#,##0 "FCFA"'
        c_pat.alignment = align_right

        for c in range(1, 13):
            ws.cell(row=current_row, column=c).border = border_fine

        current_row += 1

    # Ligne des totaux
    ws.cell(row=current_row, column=1, value="TOTAL CONSOLIDÉ").font = font_bold
    ws.cell(row=current_row, column=1).alignment = align_center
    ws.cell(row=current_row, column=4, value=f"{totaux['nb_actes']} actes").font = font_bold

    tot_brut = ws.cell(row=current_row, column=9, value=float(totaux["total_brut"]))
    tot_brut.number_format = '#,##0 "FCFA"'
    tot_brut.font = font_bold
    tot_brut.alignment = align_right

    tot_ass = ws.cell(row=current_row, column=11, value=float(totaux["total_part_assurance"]))
    tot_ass.number_format = '#,##0 "FCFA"'
    tot_ass.font = font_bold
    tot_ass.alignment = align_right

    tot_pat = ws.cell(row=current_row, column=12, value=float(totaux["total_part_patient"]))
    tot_pat.number_format = '#,##0 "FCFA"'
    tot_pat.font = font_bold
    tot_pat.alignment = align_right

    for c in range(1, 13):
        cell = ws.cell(row=current_row, column=c)
        cell.fill = fill_total
        cell.border = border_fine

    # Auto-ajustement des largeurs de colonnes
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    nom_fichier = f"bordereau_teletransmission_{annee}_{mois or 'annuel'}.xlsx"
    reponse = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    reponse["Content-Disposition"] = f'attachment; filename="{nom_fichier}"'
    classeur.save(reponse)
    return reponse
